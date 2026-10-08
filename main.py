from pathlib import Path
import openpyxl
import sqlite3

pasta = Path(__file__).parent
arquivo_excel = pasta / 'vendas_de_produtos.xlsx'

workbook = openpyxl.load_workbook(arquivo_excel)
vendas_sheet = workbook['vendas']
quantidade_registros = vendas_sheet.max_row - 1

print('=' * 50)
print('DADOS DA PLANILHA'.center(50))
print('=' * 50)
print(f'Registros encontrados: {quantidade_registros}')

banco = sqlite3.connect('vendas.db')
cursor = banco.cursor()

cursor.execute('''
    CREATE TABLE IF NOT EXISTS vendas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        linha_excel INTEGER UNIQUE,
        cliente TEXT,
        produto TEXT,
        quantidade INTEGER,
        categoria TEXT
    )
''')

for linha in vendas_sheet.iter_rows(min_row=2):
    linha_excel = linha[0].row
    cliente = linha[0].value
    produto = linha[1].value
    quantidade = linha[2].value
    categoria = linha[3].value

    cursor.execute('''
        INSERT INTO vendas (
            linha_excel,
            cliente,
            produto,
            quantidade,
            categoria
        )
         VALUES (?, ?, ?, ?, ?)

         ON CONFLICT(linha_excel) DO UPDATE SET
            cliente = excluded.cliente,
            produto = excluded.produto,
            quantidade = excluded.quantidade,
            categoria = excluded.categoria
    ''', (
        linha_excel,
        cliente,
        produto,
        quantidade,
        categoria
    ))

banco.commit()

print('\nImportação concluída com sucesso!')
print(f'{quantidade_registros} registros processados.')
print('Planilha importada no banco de dados com sucesso!')

banco.close()
